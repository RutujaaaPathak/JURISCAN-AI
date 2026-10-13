from tests.pdf_factory import image_only_pdf, multi_page_pdf

SAMPLE_CONTRACT = """1. Indemnity
The Supplier shall indemnify, defend and hold harmless the Customer against any and all losses, claims and damages of whatsoever nature.

2. Termination
The Customer may terminate this Agreement for convenience at any time on seven (7) days written notice.

3. Payment Terms
Payment shall be due within one hundred and twenty (120) days of a valid invoice.

4. Confidentiality
Each party shall keep the other party's Confidential Information confidential using reasonable care.
"""


def _upload_txt(client, content=SAMPLE_CONTRACT, filename="contract.txt", title=None):
    data = {"title": title} if title else None
    return client.post(
        "/api/documents",
        files={"file": (filename, content.encode("utf-8"), "text/plain")},
        data=data,
    )


def test_upload_txt_creates_document_analysis_and_findings(client):
    resp = _upload_txt(client, title="Sample Services Agreement")
    assert resp.status_code == 201, resp.text
    body = resp.json()

    assert body["title"] == "Sample Services Agreement"
    assert body["file_type"] == "txt"
    assert body["status"] == "PROCESSED"
    assert body["word_count"] > 0
    assert body["risk_score"] is not None
    assert body["analysis"] is not None
    assert body["analysis"]["status"] == "COMPLETED"
    assert body["analysis"]["model_name"] == "rule-based-baseline-v1"

    assert len(body["clauses"]) == 4
    findings = body["analysis"]["findings"]
    assert findings, "expected at least one risk finding"
    severities = {finding["severity"] for finding in findings}
    assert "CRITICAL" in severities
    assert "HIGH" in severities

    for finding in findings:
        assert finding["description"]
        assert finding["rule_code"]
        assert finding["clause_number"] is not None


def test_upload_docx_is_extracted(client):
    import io

    from docx import Document as DocxDocument

    document = DocxDocument()
    document.add_paragraph("1. Indemnity")
    document.add_paragraph(
        "The Supplier shall indemnify the Customer against all losses without any cap."
    )
    document.add_paragraph("2. Termination")
    document.add_paragraph(
        "Either party may terminate this Agreement for convenience at any time."
    )
    buffer = io.BytesIO()
    document.save(buffer)

    resp = client.post(
        "/api/documents",
        files={
            "file": (
                "contract.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["file_type"] == "docx"
    assert body["status"] == "PROCESSED"
    assert body["page_count"] is None
    assert body["analysis"]["findings"]


def test_duplicate_upload_returns_conflict(client):
    assert _upload_txt(client).status_code == 201
    duplicate = _upload_txt(client)
    assert duplicate.status_code == 409
    body = duplicate.json()
    assert body["error"]["code"] == "duplicate_document"
    assert body["error"]["details"]["document_id"] == 1


def test_unsupported_file_type_rejected(client):
    resp = client.post(
        "/api/documents",
        files={"file": ("notes.md", b"# heading\n\nsome text here", "text/markdown")},
    )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "unsupported_file_type"


def test_invalid_pdf_rejected(client):
    resp = client.post(
        "/api/documents",
        files={"file": ("broken.pdf", b"this is not really a pdf file", "application/pdf")},
    )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "invalid_document"


def test_empty_file_rejected(client):
    resp = client.post(
        "/api/documents",
        files={"file": ("empty.txt", b"", "text/plain")},
    )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "empty_file"


def test_list_documents_returns_uploaded_document(client):
    _upload_txt(client, title="Alpha Contract")
    resp = client.get("/api/documents")
    assert resp.status_code == 200
    body = resp.json()
    assert body["total"] == 1
    item = body["items"][0]
    assert item["title"] == "Alpha Contract"
    assert item["risk_level"] in {"LOW", "MEDIUM", "HIGH", "CRITICAL"}
    assert item["finding_count"] > 0


def test_get_document_detail_and_missing(client):
    created = _upload_txt(client).json()
    document_id = created["id"]

    resp = client.get(f"/api/documents/{document_id}")
    assert resp.status_code == 200
    assert resp.json()["id"] == document_id

    missing = client.get("/api/documents/9999")
    assert missing.status_code == 404
    assert missing.json()["error"]["code"] == "document_not_found"


def test_overview_stats_reflect_uploads(client):
    _upload_txt(client)
    resp = client.get("/api/stats/overview")
    assert resp.status_code == 200
    body = resp.json()
    assert body["total_documents"] == 1
    assert body["processed"] == 1
    assert body["processing"] == 0
    assert body["high_risk_findings"] >= 1
    assert set(body["risk_distribution"].keys()) == {"LOW", "MEDIUM", "HIGH", "CRITICAL"}
    assert len(body["upload_trend"]) == 7


def test_upload_rejects_mime_type_mismatch(client):
    resp = client.post(
        "/api/documents",
        files={"file": ("contract.pdf", b"%PDF-1.4 fake body content", "text/plain")},
    )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "mime_type_mismatch"


def test_upload_rejects_oversized_file(client, monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "max_upload_size_mb", 0.001)
    payload = b"1. Clause\n" + b"x" * 4096
    resp = client.post(
        "/api/documents",
        files={"file": ("big.txt", payload, "text/plain")},
    )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "file_too_large"


def test_upload_rejects_zip_masquerading_as_docx(client):
    import io
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("readme.txt", "not a word document")

    resp = client.post(
        "/api/documents",
        files={
            "file": (
                "contract.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "invalid_document"


def test_upload_rejects_binary_txt(client):
    resp = client.post(
        "/api/documents",
        files={"file": ("payload.txt", b"MZ\x90\x00binary payload body", "text/plain")},
    )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "invalid_document"


def test_upload_multipage_pdf_persists_page_numbers_and_extracted_at(client):
    data = multi_page_pdf(
        [
            ["1. Indemnity", "The Supplier shall indemnify the Customer without any cap."],
            ["2. Payment Terms", "Payment shall be due within one hundred (100) days."],
        ]
    )
    resp = client.post(
        "/api/documents",
        files={"file": ("contract.pdf", data, "application/pdf")},
    )
    assert resp.status_code == 201, resp.text
    body = resp.json()

    assert body["page_count"] == 2
    assert body["extracted_at"] is not None

    page_by_number = {clause["clause_number"]: clause["page_number"] for clause in body["clauses"]}
    assert page_by_number[1] == 1
    assert page_by_number[2] == 2

    for finding in body["analysis"]["findings"]:
        if finding["clause_number"] == 1:
            assert finding["clause_page"] == 1


def test_upload_persists_extracted_at_for_txt(client):
    body = _upload_txt(client).json()
    assert body["extracted_at"] is not None
    assert body["clauses"][0]["page_number"] is None


def test_upload_rejects_encrypted_pdf(client):
    data = multi_page_pdf(
        [["1. Confidential", "This document requires a password to open."]],
        encrypt=True,
    )
    resp = client.post(
        "/api/documents",
        files={"file": ("locked.pdf", data, "application/pdf")},
    )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "encrypted_document"


def test_documents_are_scoped_to_the_owner(client):
    from app.core.database import SessionLocal
    from app.models import Analysis, Document, User

    created = _upload_txt(client).json()
    own_analysis_id = created["analysis"]["id"]

    with SessionLocal() as session:
        other = User(
            email="other@example.com",
            username="other-reviewer",
            full_name="Other Reviewer",
        )
        session.add(other)
        session.flush()
        foreign = Document(
            user_id=other.id,
            title="Private matter",
            filename="private.txt",
            file_type="txt",
            status="PROCESSED",
        )
        session.add(foreign)
        session.flush()
        foreign_analysis = Analysis(
            document_id=foreign.id,
            model_name="rule-based-baseline-v1",
            status="COMPLETED",
            risk_score=10.0,
        )
        session.add(foreign_analysis)
        session.commit()
        foreign_document_id = foreign.id
        foreign_analysis_id = foreign_analysis.id

    listing = client.get("/api/documents").json()
    ids = {item["id"] for item in listing["items"]}
    assert created["id"] in ids
    assert foreign_document_id not in ids

    assert client.get(f"/api/documents/{foreign_document_id}").status_code == 404
    assert client.get(f"/api/analyses/{foreign_analysis_id}").status_code == 404
    assert client.get(f"/api/documents/{created['id']}").status_code == 200
    assert client.get(f"/api/analyses/{own_analysis_id}").status_code == 200


def test_upload_scanned_pdf_is_processed_with_ocr(client):
    import pytest

    pytest.importorskip("rapidocr_onnxruntime")
    pytest.importorskip("pymupdf")

    data = image_only_pdf(
        [
            "1. Indemnity",
            "The Supplier shall indemnify the Customer for all losses.",
        ]
    )
    resp = client.post(
        "/api/documents",
        files={"file": ("scan.pdf", data, "application/pdf")},
    )
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["status"] == "PROCESSED"
    assert body["page_count"] == 1
    assert body["extracted_at"] is not None
    assert body["analysis"]["findings"]


def test_failed_upload_is_persisted_with_error_message(client):
    resp = client.post(
        "/api/documents",
        files={"file": ("broken.pdf", b"not really a pdf body", "application/pdf")},
    )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "invalid_document"

    listing = client.get("/api/documents").json()
    assert listing["total"] == 1
    item = listing["items"][0]
    assert item["status"] == "FAILED"
    assert item["error_message"]


def test_retrying_failed_upload_reuses_the_same_row(client):
    payload = b"MZ\x90\x00binary payload pretending to be text"
    first = client.post(
        "/api/documents", files={"file": ("retry.txt", payload, "text/plain")}
    )
    second = client.post(
        "/api/documents", files={"file": ("retry.txt", payload, "text/plain")}
    )
    assert first.status_code == 422
    assert second.status_code == 422
    assert second.json()["error"]["code"] == "invalid_document"

    listing = client.get("/api/documents").json()
    assert listing["total"] == 1
    assert listing["items"][0]["status"] == "FAILED"

    overview = client.get("/api/stats/overview").json()
    assert overview["total_documents"] == 1


def test_completed_document_is_not_reprocessed(client):
    created = _upload_txt(client).json()
    analysis_id = created["analysis"]["id"]

    duplicate = _upload_txt(client)
    assert duplicate.status_code == 409
    assert duplicate.json()["error"]["code"] == "duplicate_document"

    detail = client.get(f"/api/documents/{created['id']}").json()
    assert detail["analysis"]["id"] == analysis_id
    assert client.get("/api/documents").json()["total"] == 1


def test_upload_stores_file_under_managed_upload_dir(client):
    from pathlib import Path

    from app.config import settings
    from app.core.database import SessionLocal
    from app.models import Document

    created = _upload_txt(client).json()

    with SessionLocal() as session:
        stored = session.get(Document, created["id"]).storage_path

    assert stored is not None
    stored_path = Path(stored)
    upload_dir = (Path(settings.data_dir) / "uploads").resolve()
    assert stored_path.resolve().is_relative_to(upload_dir)
    assert stored_path.exists()

    detail = client.get(f"/api/documents/{created['id']}")
    assert "storage_path" not in detail.text


def test_upload_sanitizes_a_traversal_filename(client):
    from pathlib import Path

    from app.core.database import SessionLocal
    from app.models import Document

    resp = client.post(
        "/api/documents",
        files={"file": ("../../evil.txt", SAMPLE_CONTRACT.encode("utf-8"), "text/plain")},
    )
    assert resp.status_code == 201, resp.text

    with SessionLocal() as session:
        stored = session.get(Document, resp.json()["id"]).storage_path

    basename = Path(stored).name
    assert basename.endswith("evil.txt")
    assert ".." not in basename
    assert "/" not in basename and "\\" not in basename


def test_detail_exposes_extracted_text_for_txt_without_storage_path(client):
    created = _upload_txt(client).json()
    detail = client.get(f"/api/documents/{created['id']}").json()

    assert detail["text_content"]
    assert "1. Indemnity" in detail["text_content"]
    assert detail["pages"] == []
    # Private storage locations must never leak through the API.
    assert "storage_path" not in detail


def test_detail_exposes_page_spans_for_multi_page_pdf(client):
    data = multi_page_pdf(
        [
            ["1. Indemnity", "The Supplier shall indemnify the Customer without any cap."],
            ["2. Payment Terms", "Payment shall be due within one hundred (100) days."],
        ]
    )
    created = client.post(
        "/api/documents",
        files={"file": ("contract.pdf", data, "application/pdf")},
    )
    assert created.status_code == 201, created.text
    detail = client.get(f"/api/documents/{created.json()['id']}").json()

    text = detail["text_content"]
    assert text
    assert detail["page_count"] == 2
    pages = detail["pages"]
    assert [page["page_number"] for page in pages] == [1, 2]
    assert pages[0]["start_offset"] == 0
    assert pages[0]["end_offset"] == pages[1]["start_offset"]
    assert pages[-1]["end_offset"] == len(text)
    assert all(
        page["start_offset"] < page["end_offset"] for page in pages
    )


def test_failed_document_detail_has_no_text_pages_or_storage_path(client):
    resp = client.post(
        "/api/documents",
        files={"file": ("broken.pdf", b"not really a pdf body", "application/pdf")},
    )
    assert resp.status_code == 422

    listing = client.get("/api/documents").json()
    document_id = listing["items"][0]["id"]
    detail = client.get(f"/api/documents/{document_id}").json()

    assert detail["status"] == "FAILED"
    assert detail["text_content"] is None
    assert detail["pages"] == []
    assert detail["error_message"]
    assert "storage_path" not in detail


def test_list_summaries_do_not_include_extracted_text(client):
    _upload_txt(client)
    listing = client.get("/api/documents").json()
    assert listing["items"]
    assert "text_content" not in listing["items"][0]
    assert "pages" not in listing["items"][0]


SUBCLAUSE_CONTRACT = """1. Indemnity
1.1 The Supplier shall indemnify the Customer against all losses without any cap.
1.2 This indemnity does not extend to indirect or consequential losses.

2. Payment
2.1 Payment shall be due within ninety (90) days of a valid invoice.
2.2 Interest accrues on late payment at the agreed contractual rate.
"""


def test_upload_persists_subclauses_with_headings_and_findings(client):
    body = _upload_txt(client, content=SUBCLAUSE_CONTRACT, filename="subclauses.txt").json()

    assert body["status"] == "PROCESSED"
    assert len(body["clauses"]) == 6

    headings = [clause["heading"] for clause in body["clauses"]]
    assert headings[0].startswith("1.")
    assert headings[1].startswith("1.1")
    assert headings[2].startswith("1.2")
    assert headings[3].startswith("2.")
    assert headings[4].startswith("2.1")
    assert headings[5].startswith("2.2")

    # The viewer still receives the full text and no page metadata for TXT.
    assert "1.1 The Supplier shall indemnify" in body["text_content"]
    assert body["pages"] == []
    assert all(clause["page_number"] is None for clause in body["clauses"])

    findings = body["analysis"]["findings"]
    assert findings
    clause_numbers = {clause["clause_number"] for clause in body["clauses"]}
    assert all(finding["clause_number"] in clause_numbers for finding in findings)
    assert any(finding["rule_code"] == "UNCAPPED_INDEMNITY" for finding in findings)


def test_analysis_endpoint_returns_subclause_headings(client):
    created = _upload_txt(client, content=SUBCLAUSE_CONTRACT, filename="subclauses2.txt").json()
    analysis_id = created["analysis"]["id"]

    analysis = client.get(f"/api/analyses/{analysis_id}").json()
    assert analysis["findings"]
    assert all(finding["clause_heading"] for finding in analysis["findings"])


def test_scanned_pdf_detail_exposes_ocr_text_and_page(client):
    import pytest

    pytest.importorskip("rapidocr_onnxruntime")
    pytest.importorskip("pymupdf")

    data = image_only_pdf(
        [
            "1. Indemnity",
            "The Supplier shall indemnify the Customer for all losses.",
        ]
    )
    created = client.post(
        "/api/documents",
        files={"file": ("scan.pdf", data, "application/pdf")},
    )
    assert created.status_code == 201, created.text
    detail = client.get(f"/api/documents/{created.json()['id']}").json()

    assert detail["status"] == "PROCESSED"
    assert detail["text_content"]
    assert "indemn" in detail["text_content"].lower()
    assert [page["page_number"] for page in detail["pages"]] == [1]


def test_download_analysis_report_is_a_real_pdf_with_analysis_data(client):
    from io import BytesIO

    from pypdf import PdfReader

    created = _upload_txt(client, title="Report Test Contract").json()
    report = client.get(f"/api/documents/{created['id']}/report")

    assert report.status_code == 200
    assert report.headers["content-type"] == "application/pdf"
    assert (
        report.headers.get("content-disposition", "").startswith("attachment;")
    )
    assert report.content.startswith(b"%PDF")

    reader = PdfReader(BytesIO(report.content))
    assert len(reader.pages) >= 1
    text = "\n".join(page.extract_text() or "" for page in reader.pages)

    assert "Report Test Contract" in text
    assert "Legal safety" in text
    assert "Overall assessment" in text
    assert "Issues to review" in text
    assert "UNCAPPED_INDEMNITY" in text or "indemnify" in text.lower()
    assert "not a legal opinion" in text


def test_download_analysis_report_not_available_before_processing(client):
    resp = client.post(
        "/api/documents",
        files={"file": ("broken.pdf", b"not really a pdf body", "application/pdf")},
    )
    assert resp.status_code == 422

    listing = client.get("/api/documents").json()
    document_id = listing["items"][0]["id"]
    report = client.get(f"/api/documents/{document_id}/report")

    assert report.status_code == 422
    assert report.json()["error"]["code"] == "report_not_available"


def test_delete_document_removes_row_results_and_stored_file(client):
    from pathlib import Path

    from sqlalchemy import func, select

    from app.core.database import SessionLocal
    from app.models import Analysis, Clause, Document, RiskFinding

    created = _upload_txt(client, title="Delete Me", filename="delete-me.txt").json()
    document_id = created["id"]

    with SessionLocal() as db:
        document = db.get(Document, document_id)
        assert document is not None
        stored_path = Path(document.storage_path)
        assert stored_path.is_file()
        analysis_ids = [analysis.id for analysis in document.analyses]
        assert analysis_ids

    resp = client.delete(f"/api/documents/{document_id}")
    assert resp.status_code == 204, resp.text
    assert resp.content == b""

    assert client.get(f"/api/documents/{document_id}").status_code == 404
    listing = client.get("/api/documents").json()
    assert all(item["id"] != document_id for item in listing["items"])
    assert not stored_path.exists()

    with SessionLocal() as db:
        assert db.get(Document, document_id) is None
        assert (
            db.scalar(
                select(func.count())
                .select_from(Analysis)
                .where(Analysis.document_id == document_id)
            )
            == 0
        )
        assert (
            db.scalar(
                select(func.count())
                .select_from(Clause)
                .where(Clause.document_id == document_id)
            )
            == 0
        )
        assert (
            db.scalar(
                select(func.count())
                .select_from(RiskFinding)
                .where(RiskFinding.analysis_id.in_(analysis_ids))
            )
            == 0
        )


def test_delete_missing_document_returns_404(client):
    resp = client.delete("/api/documents/999999")

    assert resp.status_code == 404
    assert resp.json()["error"]["code"] == "document_not_found"


def test_delete_document_updates_overview_stats(client):
    created = _upload_txt(client, title="Stats Delete", filename="stats-delete.txt").json()

    before = client.get("/api/stats/overview").json()
    assert before["total_documents"] == 1
    assert before["processed"] == 1

    assert client.delete(f"/api/documents/{created['id']}").status_code == 204

    after = client.get("/api/stats/overview").json()
    assert after["total_documents"] == 0
    assert after["processed"] == 0
    assert after["high_risk_findings"] == 0
    assert after["category_counts"] == []


def test_delete_all_documents_removes_every_document_and_results(client):
    from pathlib import Path

    from sqlalchemy import func, select

    from app.core.database import SessionLocal
    from app.models import Analysis, Clause, Document, RiskFinding

    first = _upload_txt(client, title="Bulk A", filename="bulk-a.txt").json()
    second = _upload_txt(
        client,
        content=SAMPLE_CONTRACT + "\n5. Assignment\nThe Supplier may not assign this Agreement.",
        title="Bulk B",
        filename="bulk-b.txt",
    ).json()

    with SessionLocal() as db:
        paths = [
            Path(db.get(Document, first["id"]).storage_path),
            Path(db.get(Document, second["id"]).storage_path),
        ]
        assert all(path.is_file() for path in paths)

    resp = client.delete("/api/documents")
    assert resp.status_code == 200, resp.text
    assert resp.json() == {"deleted": 2}

    listing = client.get("/api/documents").json()
    assert listing["items"] == []
    assert listing["total"] == 0
    assert all(not path.exists() for path in paths)

    stats = client.get("/api/stats/overview").json()
    assert stats["total_documents"] == 0
    assert stats["category_counts"] == []

    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(Document)) == 0
        assert db.scalar(select(func.count()).select_from(Analysis)) == 0
        assert db.scalar(select(func.count()).select_from(Clause)) == 0
        assert db.scalar(select(func.count()).select_from(RiskFinding)) == 0


def test_delete_all_on_empty_database_returns_zero(client):
    resp = client.delete("/api/documents")

    assert resp.status_code == 200
    assert resp.json() == {"deleted": 0}


def test_delete_all_preserves_other_users_documents(client):
    from app.core.database import SessionLocal
    from app.models import Document, User

    created = _upload_txt(client, title="Mine", filename="mine.txt").json()

    with SessionLocal() as db:
        other = User(email="other@example.com", username="other", full_name="Other")
        db.add(other)
        db.flush()
        foreign = Document(
            user_id=other.id,
            title="Theirs",
            filename="theirs.txt",
            file_type="txt",
            status="PROCESSED",
            content_hash="foreign-hash",
        )
        db.add(foreign)
        db.commit()
        foreign_id = foreign.id

    resp = client.delete("/api/documents")
    assert resp.status_code == 200
    assert resp.json() == {"deleted": 1}

    with SessionLocal() as db:
        assert db.get(Document, created["id"]) is None
        assert db.get(Document, foreign_id) is not None
