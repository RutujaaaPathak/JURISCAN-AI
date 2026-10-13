"""Text extraction and upload validation for PDF, DOCX and TXT files.

Native text is preferred. For PDFs, pages with little or no extractable text
(scanned / image-only pages) are optionally run through OCR; native text is
never discarded in favour of a shorter OCR result. When OCR is unavailable or
disabled and a document is entirely image-only, a clear ``ocr_unavailable``
error is raised rather than silently returning empty text.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone

from app.config import settings
from app.core.exceptions import ValidationFailed
from app.services import ocr

SUPPORTED_EXTENSIONS = ("pdf", "docx", "txt")

MIME_BY_EXTENSION = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
}

# Declared MIME types that are accepted for each extension. Browsers and API
# clients are inconsistent, so a small set of official and generic values is
# allowed per type; anything else (e.g. a PDF declared as text/plain) is
# rejected. Generic octet-stream values are always tolerated because they carry
# no reliable type information.
ALLOWED_MIME_BY_EXTENSION = {
    "pdf": {"application/pdf", "application/x-pdf"},
    "docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
        "application/zip",
        "application/x-zip-compressed",
    },
    "txt": {"text/plain"},
}

GENERIC_MIME_TYPES = {"application/octet-stream", "binary/octet-stream"}

PDF_MAGIC = b"%PDF"
ZIP_MAGIC = b"PK\x03\x04"

# Minimal magic-byte prefixes of common executable/archive formats that should
# never masquerade as a plain text document.
BINARY_MAGIC_PREFIXES = (b"MZ", b"\x7fELF", b"\xca\xfe\xba\xbe", b"\x00asm")

DOCX_CONTENT_MARKER = "word/document.xml"

MIN_TEXT_CHARACTERS = 20


@dataclass(slots=True)
class PageSpan:
    """Character range of a single source page within the extracted text.

    Offsets are inclusive of ``start_offset`` and exclusive of ``end_offset``,
    matching the convention used by ``Segment``. Page spans are only produced
    for paginated formats (currently PDF); other types report an empty list.
    """

    page_number: int
    start_offset: int
    end_offset: int
    char_count: int


@dataclass(slots=True)
class ExtractedDocument:
    text: str
    file_type: str
    page_count: int | None
    word_count: int
    language: str
    extracted_at: datetime
    pages: list[PageSpan] = field(default_factory=list)
    ocr_pages: list[int] = field(default_factory=list)


def _extension_of(filename: str) -> str:
    if "." not in filename:
        return ""
    return filename.rsplit(".", 1)[-1].strip().lower()


def _detect_language(text: str) -> str:
    """Very small heuristic language hint. Defaults to English."""

    sample = text[:2000]
    if not sample:
        return "en"
    letters = [ch for ch in sample if ch.isalpha()]
    if not letters:
        return "en"
    non_ascii = sum(1 for ch in letters if ord(ch) > 127)
    return "en" if non_ascii / len(letters) < 0.2 else "unknown"


def _clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_pdf(data: bytes) -> tuple[list[str], int]:
    """Return the raw text of each PDF page, in order, plus the page count."""

    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            # ``decrypt("")`` returns a PasswordType (NOT_DECRYPTED == 0) when
            # the document cannot be opened with an empty password; it does not
            # raise. Accessing ``.pages`` afterwards would raise a pypdf error,
            # so detect the failure explicitly to report an accurate code.
            result = reader.decrypt("")
            if reader.is_encrypted or not result:
                raise ValidationFailed(
                    "The PDF is password protected and could not be opened.",
                    code="encrypted_document",
                )
        pages = reader.pages
        chunks = [page.extract_text() or "" for page in pages]
        return chunks, len(pages)
    except ValidationFailed:
        raise
    except Exception as exc:
        raise ValidationFailed(
            "The PDF could not be parsed. The file may be corrupt.",
            code="invalid_document",
        ) from exc


def _apply_ocr(
    data: bytes,
    chunks: list[str],
    *,
    allow_ocr: bool,
    min_page_chars: int,
    dpi: int,
) -> tuple[list[str], list[int]]:
    """OCR only the PDF pages that carry little or no native text.

    Native text is preserved: an OCR result replaces a page only when it is
    strictly longer than the native text found on that page. Returns the
    (possibly) updated per-page chunks and the 1-based pages that gained text
    thanks to OCR.
    """

    cleaned = [_clean_text(chunk) for chunk in chunks]
    textless = [index + 1 for index, page in enumerate(cleaned) if len(page) < min_page_chars]
    if not textless:
        return chunks, []

    engine_ready = allow_ocr and ocr.ocr_available()
    if not engine_ready:
        if len(textless) == len(chunks):
            raise ValidationFailed(
                "This PDF appears to be a scanned or image-only document and OCR "
                "is not enabled or available.",
                code="ocr_unavailable",
                details={"pages": textless},
            )
        # Only some pages are image-only; keep the native text we do have.
        return chunks, []

    try:
        readings = ocr.ocr_pdf_pages(data, textless, dpi=dpi)
    except ocr.OcrUnavailable as exc:
        if len(textless) == len(chunks):
            raise ValidationFailed(
                "This PDF appears to be a scanned or image-only document but the "
                "OCR engine is not available.",
                code="ocr_unavailable",
                details={"pages": textless},
            ) from exc
        return chunks, []
    except Exception as exc:
        raise ValidationFailed(
            "OCR failed while reading the scanned pages of the PDF.",
            code="ocr_failed",
            details={"pages": textless},
        ) from exc

    updated = list(chunks)
    ocr_pages: list[int] = []
    for number, raw in readings.items():
        ocr_text = _clean_text(raw)
        if 1 <= number <= len(cleaned) and len(ocr_text) > len(cleaned[number - 1]):
            updated[number - 1] = ocr_text
            ocr_pages.append(number)
    return updated, sorted(ocr_pages)


def _assemble_pages(pages: list[str]) -> tuple[str, list[PageSpan]]:
    """Join per-page text into one document while recording page offsets.

    Each page is normalised independently and joined with a blank line so that
    page spans map cleanly back to character offsets in the extracted text.
    """

    spans: list[PageSpan] = []
    text = ""
    for index, raw in enumerate(pages, start=1):
        page_text = _clean_text(raw)
        if index > 1:
            text += "\n\n"
        start = len(text)
        text += page_text
        spans.append(
            PageSpan(
                page_number=index,
                start_offset=start,
                end_offset=len(text),
                char_count=len(page_text),
            )
        )
    return text, spans


def page_number_for_offset(spans: list[PageSpan], offset: int | None) -> int | None:
    """Map a character offset to a 1-based page number, if known."""

    if offset is None or not spans:
        return None
    for span in spans:
        if span.start_offset <= offset < span.end_offset:
            return span.page_number
    if offset >= spans[-1].end_offset:
        return spans[-1].page_number
    return spans[0].page_number


def _extract_docx(data: bytes) -> tuple[str, int | None]:
    from docx import Document as DocxDocument

    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise ValidationFailed(
            "The DOCX file could not be parsed. The file may be corrupt.",
            code="invalid_document",
        ) from exc

    parts: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts), None


def _extract_txt(data: bytes) -> tuple[str, int | None]:
    if data.startswith(BINARY_MAGIC_PREFIXES):
        raise ValidationFailed(
            "The file appears to be a binary or executable file, not a text document.",
            code="invalid_document",
        )
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding), None
        except UnicodeError:
            continue
    raise ValidationFailed(
        "The text file could not be decoded. Use UTF-8, UTF-16 or Latin-1 encoding.",
        code="invalid_document",
    )


def _assert_magic(data: bytes, file_type: str) -> None:
    if len(data) < 4:
        raise ValidationFailed(
            "The uploaded file is empty or too small to be a valid document.",
            code="invalid_document",
        )
    if file_type == "pdf" and not data.startswith(PDF_MAGIC):
        raise ValidationFailed(
            "The file does not look like a valid PDF (missing %PDF header).",
            code="invalid_document",
        )
    if file_type == "docx" and not data.startswith(ZIP_MAGIC):
        raise ValidationFailed(
            "The file does not look like a valid DOCX archive.",
            code="invalid_document",
        )


def _assert_mime(declared_content_type: str | None, file_type: str) -> None:
    """Reject a declared MIME type that contradicts the file extension.

    Generic ``application/octet-stream`` values are accepted because they carry
    no reliable type information; the magic-byte and parse checks below remain
    the authoritative validation.
    """

    if not declared_content_type:
        return
    normalized = declared_content_type.split(";", 1)[0].strip().lower()
    if not normalized or normalized in GENERIC_MIME_TYPES:
        return
    if normalized in ALLOWED_MIME_BY_EXTENSION.get(file_type, set()):
        return
    if file_type == "txt" and normalized.startswith("text/"):
        return
    raise ValidationFailed(
        f"The declared content type '{normalized}' does not match the '{file_type}' extension.",
        code="mime_type_mismatch",
        details={
            "declared": normalized,
            "expected": MIME_BY_EXTENSION[file_type],
        },
    )


def _assert_docx_integrity(data: bytes) -> None:
    """Verify the archive really is a Word document, not just any ZIP file."""

    import zipfile

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
    except Exception as exc:
        raise ValidationFailed(
            "The DOCX file could not be opened. The archive may be corrupt.",
            code="invalid_document",
        ) from exc
    if DOCX_CONTENT_MARKER not in names:
        raise ValidationFailed(
            "The file is a ZIP archive but is not a valid Word document.",
            code="invalid_document",
        )


def resolve_file_type(
    filename: str, declared_content_type: str | None = None
) -> str:
    """Validate the extension/MIME pair and return the resolved file type."""

    file_type = _extension_of(filename)
    if file_type not in SUPPORTED_EXTENSIONS:
        raise ValidationFailed(
            "Unsupported file type. Upload a PDF, DOCX or TXT file.",
            code="unsupported_file_type",
            details={"filename": filename, "supported": list(SUPPORTED_EXTENSIONS)},
        )
    _assert_mime(declared_content_type, file_type)
    return file_type


def validate_and_extract(
    data: bytes,
    filename: str,
    declared_content_type: str | None = None,
    max_size_bytes: int | None = None,
    *,
    allow_ocr: bool | None = None,
    ocr_min_page_chars: int | None = None,
    ocr_dpi: int | None = None,
) -> ExtractedDocument:
    """Validate an upload and extract its text.

    Raises ValidationFailed with a stable error code when the upload is not a
    supported or parseable document. ``allow_ocr``/``ocr_min_page_chars``/
    ``ocr_dpi`` override the corresponding settings (used by tests).
    """

    if max_size_bytes is not None and len(data) > max_size_bytes:
        raise ValidationFailed(
            f"The file exceeds the maximum size of {max_size_bytes // (1024 * 1024)} MB.",
            code="file_too_large",
        )

    file_type = resolve_file_type(filename, declared_content_type)
    _assert_magic(data, file_type)

    ocr_pages: list[int] = []
    if file_type == "pdf":
        chunks, page_count = _extract_pdf(data)
        chunks, ocr_pages = _apply_ocr(
            data,
            chunks,
            allow_ocr=settings.ocr_enabled if allow_ocr is None else allow_ocr,
            min_page_chars=(
                settings.ocr_min_page_chars
                if ocr_min_page_chars is None
                else ocr_min_page_chars
            ),
            dpi=settings.ocr_render_dpi if ocr_dpi is None else ocr_dpi,
        )
        text, pages = _assemble_pages(chunks)
    elif file_type == "docx":
        _assert_docx_integrity(data)
        raw, page_count = _extract_docx(data)
        text, pages = _clean_text(raw), []
    else:
        raw, page_count = _extract_txt(data)
        text, pages = _clean_text(raw), []

    if len(text) < MIN_TEXT_CHARACTERS:
        raise ValidationFailed(
            "No extractable text was found in the document. If this is a scanned "
            "or image-only file, ensure OCR is enabled and available.",
            code="no_extractable_text",
        )

    word_count = len(text.split())
    return ExtractedDocument(
        text=text,
        file_type=file_type,
        page_count=page_count,
        word_count=word_count,
        language=_detect_language(text),
        extracted_at=datetime.now(timezone.utc),
        pages=pages,
        ocr_pages=ocr_pages,
    )


__all__ = [
    "ALLOWED_MIME_BY_EXTENSION",
    "ExtractedDocument",
    "MIME_BY_EXTENSION",
    "PageSpan",
    "SUPPORTED_EXTENSIONS",
    "page_number_for_offset",
    "resolve_file_type",
    "validate_and_extract",
]
